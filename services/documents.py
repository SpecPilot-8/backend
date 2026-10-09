from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
import pymupdf

from api.errors import APIError

MAX_TEXT = 1_000_000
MAX_CELLS = 200_000


def _check_archive(data: bytes):
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise APIError(413, "EXPANDED_DOCUMENT_TOO_LARGE", "압축 해제된 문서가 100MB 제한을 초과합니다")
    except BadZipFile as error:
        raise APIError(422, "INVALID_DOCUMENT", "유효한 Office 문서가 아닙니다") from error


def extract(data: bytes, filename: str) -> tuple[str, list[dict], list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in (".docx", ".pdf", ".xlsx"):
        raise APIError(415, "UNSUPPORTED_DOCUMENT_FORMAT", "지원하는 파일 형식은 DOCX·PDF·XLSX입니다")
    if not data:
        raise APIError(422, "EMPTY_DOCUMENT", "빈 파일은 등록할 수 없습니다")
    blocks, warnings = [], []
    text_size = 0

    def append(block: dict):
        nonlocal text_size
        text_size += len(block["text"])
        if text_size > MAX_TEXT:
            raise APIError(413, "EXTRACTED_TEXT_TOO_LARGE", "추출 텍스트가 100만자 제한을 초과합니다")
        block["block_id"] = f"B-{len(blocks) + 1}"
        blocks.append(block)

    try:
        if suffix == ".docx":
            _check_archive(data)
            document = WordDocument(BytesIO(data))
            for index, element in enumerate(document.element.body):
                if element.tag == qn("w:p"):
                    text = Paragraph(element, document).text
                    if text.strip():
                        append({"kind": "paragraph", "text": text, "location": {"body_index": index}})
                elif element.tag == qn("w:tbl"):
                    rows = [[cell.text for cell in row.cells] for row in Table(element, document).rows]
                    append({"kind": "table", "text": "\n".join(" | ".join(r) for r in rows),
                            "rows": rows, "location": {"body_index": index}})
            warnings.append("DOCX 텍스트박스·도형·중첩 표의 별도 구조 및 렌더링 페이지 번호는 추출하지 않습니다")
        elif suffix == ".pdf":
            with pymupdf.open(stream=data, filetype="pdf") as document:
                if document.needs_pass:
                    raise APIError(422, "ENCRYPTED_DOCUMENT", "암호화된 PDF는 지원하지 않습니다")
                for page in document:
                    if page.number >= 1000:
                        raise APIError(413, "TOO_MANY_PAGES", "PDF는 최대 1000페이지까지 지원합니다")
                    page_has_text = False
                    for block in page.get_text("blocks", sort=True):
                        if block[6] == 0 and block[4].strip():
                            append({"kind": "text", "text": block[4],
                                    "location": {"page": page.number + 1, "bbox": list(block[:4])}})
                            page_has_text = True
                    if not page_has_text:
                        warnings.append(f"PDF {page.number + 1}페이지에 추출 가능한 텍스트가 없습니다. OCR이 필요합니다")
            warnings.append("PDF 표의 셀 구조 복원과 이미지·도형 OCR은 아직 연결되지 않았습니다")
        else:
            _check_archive(data)
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                visited = 0
                for sheet in workbook:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                        raise APIError(413, "TOO_MANY_CELLS", "Excel 시트 크기가 셀 수 제한을 초과합니다")
                    for row in sheet.iter_rows():
                        for cell in row:
                            visited += 1
                            if visited > MAX_CELLS:
                                raise APIError(413, "TOO_MANY_CELLS", "Excel은 최대 20만 셀까지 지원합니다")
                            if cell.value is not None:
                                append({"kind": "cell", "text": str(cell.value),
                                        "location": {"sheet": sheet.title, "cell": cell.coordinate}})
            finally:
                workbook.close()
            warnings.append("Excel 도형 안 텍스트는 추출하지 않으며 수식은 계산하지 않고 원문으로 보존합니다")
    except APIError:
        raise
    except Exception as error:
        raise APIError(422, "INVALID_DOCUMENT", "문서를 읽을 수 없습니다. 파일 내용과 형식을 확인하세요") from error
    if not blocks:
        warnings.append("추출된 텍스트가 없습니다. 스캔 문서·도형 문서인지 확인하세요")
    return "\n\n".join(block["text"] for block in blocks), blocks, warnings
