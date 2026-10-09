from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def _workbook(title, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
        # Explicit string cells prevent user-supplied =... content becoming executable Excel formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(color="FFFFFF", bold=True)
        sheet.column_dimensions[cell.column_letter].width = 25
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return workbook, sheet


def test_spec(cases, reqs, mismatch_notes=None):
    mismatch_notes = mismatch_notes or {}
    headers = ["Test ID", "요구사항 ID", "시나리오명", "유형", "사전조건", "시험절차", "입력값",
               "기대결과", "실제결과", "판정", "비고", "데이터 출처", "요구사항 버전"]
    rows = [[c.test_id, reqs[c.requirement_id].req_id, c.title, c.type, c.preconditions,
             "\n".join(f"{i}. {step}" for i, step in enumerate(c.steps, start=1)), c.input_data,
             c.expected, c.actual, c.result,
             "\n".join(filter(None, [c.notes, mismatch_notes.get(c.requirement_id)])), c.source, c.revision]
            for c in cases]
    workbook, sheet = _workbook("테스트 사양서", headers, rows)
    dropdown = DataValidation(type="list", formula1='"PASS,FAIL"', allow_blank=True)
    sheet.add_data_validation(dropdown)
    dropdown.add(f"J2:J{len(cases) + 1}")
    return _bytes(workbook)


def traceability(project, snapshot, records):
    headers = ["요구사항 ID", "기능명", "조건", "처리", "기대결과", "구현상태", "판정사유",
               "근거 코드", "판정 출처", "요구사항 버전", "스냅샷 ID"]
    rows = [[req.req_id, req.title, req.condition, req.action, req.expected,
             result["status"] if result else "not_run", result["reason"] if result else "",
             "\n".join(f"{e['file_path']}:{e['start_line']}-{e['end_line']}" for e in result["evidence"])
             if result else "", result["source"] if result else "", req.revision,
             snapshot.id if snapshot else ""] for req, result in records]
    workbook, _ = _workbook("요구사항 추적표", headers, rows)
    info = workbook.create_sheet("실행 정보")
    info.append(["프로젝트", project.name])
    info.append(["스냅샷", snapshot.id if snapshot else "not_indexed"])
    info.append(["판정 상태", "AI 미연결. 수동 등록 판정만 표시하며 빈 판정은 not_run입니다"])
    for row in info:
        for cell in row:
            cell.data_type = "s"
    return _bytes(workbook)


def _bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
