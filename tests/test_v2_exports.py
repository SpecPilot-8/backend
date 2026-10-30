from copy import deepcopy
from io import BytesIO
from uuid import uuid4
from openpyxl import load_workbook


def base(p):
    return f'/api/v2/projects/{p}'


def export_body(case, revision=2, detail=False):
    return {'test_specification_id': case['test_specification_id'], 'revision': revision,
            'execution_selection': 'latest_per_case', 'template_id': 'thinktree-test-spec-v1',
            'include_analysis_detail': detail, 'format': 'xlsx'}


def workbook(client, p, job):
    assert job['status'] == 'completed'
    r = client.get(base(p) + f'/artifacts/{job["result_ref"]["id"]}/download')
    assert r.status_code == 200 and 'attachment' in r.headers['content-disposition']
    return load_workbook(BytesIO(r.content)), r.content


def test_test_export_steps_actual_results_and_literal_text(client, v2_project, v2_case):
    steps = deepcopy(v2_case['steps'])
    steps[0]['action'] = '=HYPERLINK("https://example.com", "literal")'
    steps[0]['test_data']['notes'] = '시험 입력 비고'
    changed = client.patch(base(v2_project) + f'/test-cases/{v2_case["id"]}/procedure', json={'base_revision': 1, 'steps': steps}).json()
    e = client.post(base(v2_project) + f'/test-cases/{v2_case["id"]}/executions', json={'test_case_revision': 2, 'client_execution_id': str(uuid4())}).json()
    r = client.patch(base(v2_project) + f'/test-executions/{e["id"]}', json={'base_version': 1,'client_edit_id':str(uuid4()),'step_results':[
        {'step_key':'step-1','status':'failed','actual':'실제 실패','notes':'수행 메모','performed_at':'2026-10-11T05:00:00Z'}]})
    assert r.status_code == 200
    job = client.post(base(v2_project) + '/exports/test-specification', json=export_body(v2_case,3,True)).json()
    w, _ = workbook(client,v2_project,job)
    assert w.sheetnames == ['개정이력','00. 목차','04. 공지사항','요구사항·시험 대응','시험 설계 점검','분석 상세']
    assert w['시험 설계 점검']['F2'].value == '관점 연결 없음'
    assert w['분석 상세']['R2'].value == '설계 사유 미기록'
    s = w['04. 공지사항']
    assert s['A7'].value == s['A8'].value == v2_case['test_no']
    assert s['D7'].value == 1 and s['D8'].value == 2
    assert s['E7'].value == steps[0]['action'] and s['E7'].data_type == 's'
    assert s['F7'].value == steps[0]['expected']
    assert s['I7'].value == '시험 입력 비고'
    assert s['J7'].value == 'X' and s['J8'].value is None
    assert s['J2'].value is not None and s['H1'].value is None and s['J1'].value is None
    assert 'E4:E5' in {str(x) for x in s.merged_cells.ranges}
    assert w['분석 상세']['K2'].value == '수행 메모'
    assert not any(c.data_type == 'f' for sheet in w for row in sheet for c in row)
    assert changed['prediction']['status'] == 'unknown'


def test_export_pins_revision_and_execution_and_scopes_download(client, v2_project, v2_case):
    job = client.post(base(v2_project) + '/exports/test-specification', json=export_body(v2_case)).json()
    first, first_bytes = workbook(client,v2_project,job)
    client.patch(base(v2_project) + f'/test-cases/{v2_case["id"]}/procedure', json={'base_revision':1,'preconditions':'새 절차'})
    second, second_bytes = workbook(client,v2_project,job)
    assert first_bytes == second_bytes
    assert first['04. 공지사항']['J7'].value is None
    old_job = client.post(base(v2_project) + '/exports/test-specification', json=export_body(v2_case)).json()
    old,_ = workbook(client,v2_project,old_job)
    assert old['04. 공지사항']['D8'].value == 2
    other = client.post('/api/v2/projects',json={'name':'다른 프로젝트'}).json()['id']
    assert client.get(base(other) + f'/artifacts/{job["result_ref"]["id"]}/download').status_code == 404


def test_empty_export_rejected_and_trace_includes_unanalyzed_state(client,v2_project,v2_specification,v2_run):
    r = client.post(base(v2_project) + '/exports/test-specification', json={'test_specification_id':v2_specification['id'], 'revision':1,'execution_selection':'latest_per_case','template_id':'thinktree-test-spec-v1','format':'xlsx'})
    assert r.status_code == 409 and r.json()['error']['code'] == 'EMPTY_EXPORT'
    job = client.post(base(v2_project) + '/exports/traceability',json={'verification_run_id':v2_run['id'],'format':'xlsx'}).json()
    w,_ = workbook(client,v2_project,job)
    assert w['요구사항 추적표']['G2'].value == '미분석'
    assert w['수락조건']['H2'].value == '미분석'
    assert w['코드 근거'].max_row == 1


def test_export_limit_and_text_conversion(client, config, v2_project, v2_case):
    from api.v2.exports import value
    from openpyxl import Workbook
    from api.errors import APIError
    import pytest

    object.__setattr__(config, 'max_export_rows', 1)
    r = client.post(base(v2_project) + '/exports/test-specification', json=export_body(v2_case))
    assert r.status_code == 413 and r.json()['error']['code'] == 'EXPORT_LIMIT_EXCEEDED'
    sheet, warnings = Workbook().active, []
    value(sheet['A1'], '가' * 33000, warnings)
    assert len(sheet['A1'].value) <= 32767 and warnings[0]['code'] == 'EXCEL_TEXT_TRUNCATED'
    with pytest.raises(APIError) as error:
        value(sheet['A2'], '불가\x00문자', warnings)
    assert error.value.code == 'UNEXPORTABLE_TEXT'


def test_design_review_sheet_obeys_row_limit(client, config, v2_project, v2_case):
    # 단계 수는 한도 이하지만 한 조건에서 파생된 여러 시험 관점은 한도를 넘는다.
    object.__setattr__(config, 'max_export_rows', 2)
    assert client.post(base(v2_project) + '/exports/test-specification', json=export_body(v2_case)).status_code == 202
    r = client.post(base(v2_project) + '/exports/test-specification', json=export_body(v2_case, detail=True))
    assert r.status_code == 413 and r.json()['error']['code'] == 'EXPORT_LIMIT_EXCEEDED'


def test_export_includes_original_conditions_and_pins_manual_review(client, v2_project, v2_case):
    from test_v2_test_design_reviews import body, endpoint
    url = endpoint(v2_project, v2_case)
    request = body(v2_case)
    assert client.put(url, json=request).status_code == 200
    response = client.post(f'/api/v2/projects/{v2_project}/exports/test-specification', json=export_body(v2_case, detail=True))
    assert response.status_code == 202, response.text
    book, content = workbook(client, v2_project, response.json())
    sheet = book['요구사항·시험 대응']
    assert sheet['D2'].value and sheet['E2'].value
    assert sheet['I2'].value == v2_case['steps'][0]['action']
    assert sheet['J2'].value == v2_case['steps'][0]['expected']
    assert sheet['L2'].value == '설계 적합·사용자 검토'
    assert sheet['M2'].value == request['items'][0]['reason']
    assert book['04. 공지사항']['J7'].value is None
    request.update(base_version=1, client_review_id=str(uuid4()))
    request['items'][0]['decision'] = 'needs_changes'
    assert client.put(url, json=request).status_code == 200
    _, pinned = workbook(client, v2_project, response.json())
    assert pinned == content
