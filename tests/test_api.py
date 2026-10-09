from io import BytesIO

import pymupdf
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from api.main import create_app




def params(project, snapshot=None):
    result = {"project_id": project["id"]}
    if snapshot:
        result["snapshot_id"] = snapshot["id"]
    return result


def chunks(client, project, snapshot):
    return client.get("/api/v1/code/chunks", params=params(project, snapshot)).json()["items"]


def verdict(client, project, requirement, snapshot, status="mismatch", **extra):
    return client.post("/api/v1/verification-results/import", json={
        "project_id": project["id"], "requirement_id": requirement["id"], "revision": requirement["revision"],
        "snapshot_id": snapshot["id"], "status": status, "reason": "5회 잠금 요구사항이 10회로 구현됨",
        "evidence_chunk_ids": [chunks(client, project, snapshot)[0]["id"]], **extra,
    })


def case_body(project, requirement):
    return {"project_id": project["id"], "requirement_id": requirement["id"], "revision": requirement["revision"],
            "test_id": "TC-001", "title": "=SUM(1,2)", "type": "boundary", "preconditions": "실패 4회",
            "steps": ["잘못된 비밀번호로 로그인"], "input_data": "wrong", "expected": "계정 잠금"}


def docx_bytes():
    document = Document()
    document.add_paragraph("실패 5회 시 계정을 잠근다")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "REQ-001"
    table.cell(0, 1).text = "잠금"
    result = BytesIO()
    document.save(result)
    return result.getvalue()


def xlsx_bytes():
    workbook = Workbook()
    workbook.active["B2"] = "REQ-001"
    workbook.active["C2"] = "=1+1"
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def pdf_bytes(text=True):
    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((50, 50), "Lock after 5 failures")
    result = document.tobytes()
    document.close()
    return result


def test_core_project_persistence_and_routes(client, config, project):
    with TestClient(create_app(config)) as reopened:
        assert reopened.get(f"/api/v1/projects/{project['id']}").json()["name"] == "Demo"
        assert reopened.get("/health").json()["status"] == "ok"
        paths = reopened.get("/openapi.json").json()["paths"]
        assert "/api/v1/projects" in paths and "/api/v1/jobs" in paths
        assert reopened.get("/docs").status_code == 200
        assert reopened.get("/api/v1/jobs", params=params(project)).json()["total"] == 0
    invalid = client.post("/api/v1/projects", json={"name": "Outside", "workspace_path": str(config.workspace_root.parent)})
    assert invalid.status_code == 422
