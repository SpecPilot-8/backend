from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Literal

from fastapi import APIRouter
from fastapi.responses import FileResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import Field

from api.errors import APIError
from api.models import new_id
from api.v2.common import Config, DB, Input, make_job
from api.v2.storage import create, historical, project, resource
from api.v2.testing import EXPORT_SECTIONS, current_cases, latest_execution, lock_project
from api.v2.verification import selected_rows

router = APIRouter(prefix='/api/v2/projects/{p}', tags=['v2 Excel 출력·다운로드'])
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_INVALID_XML = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


class TraceExport(Input):
    verification_run_id: str
    format: Literal['xlsx']


class TestExport(Input):
    test_specification_id: str
    revision: int = Field(ge=1)
    execution_selection: Literal['latest_per_case']
    template_id: Literal['thinktree-test-spec-v1']
    include_analysis_detail: bool = False
    format: Literal['xlsx']


def value(cell, content, warnings):
    if isinstance(content, str):
        if _INVALID_XML.search(content):
            raise APIError(422, 'UNEXPORTABLE_TEXT', 'Excel에 표시할 수 없는 제어 문자가 있습니다')
        if len(content) > 32767:
            content = content[:32720] + '\n[Excel 셀 한도로 일부 표시됨]'
            warnings.append({'code': 'EXCEL_TEXT_TRUNCATED', 'message': '셀 한도 초과: API의 원문/청크 조회에서 전체 내용을 확인하세요', 'cell': f'{cell.parent.title}!{cell.coordinate}'})
        cell.value = content
        cell.data_type = 's'  # Literal text including =, +, -, @ must never become formulas.
    else:
        cell.value = content
    if isinstance(content, datetime):
        cell.number_format = 'yyyy-mm-dd'
    cell.font = Font(name='맑은 고딕', size=10)
    cell.alignment = Alignment(vertical='center', wrap_text=True)


def row(sheet, index, values, warnings):
    for col, content in enumerate(values, 1):
        value(sheet.cell(index, col), content, warnings)


def style_grid(sheet, first, last, cols):
    thin = Side(style='thin', color='C4C4C4')
    for cells in sheet.iter_rows(min_row=first, max_row=last, max_col=cols):
        for cell in cells:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f'E{first+2}'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.print_area = f'A1:{get_column_letter(cols)}{last}'


def header(sheet, at, values, warnings):
    row(sheet, at, values, warnings)
    for cell in sheet[at]:
        cell.fill = PatternFill('solid', fgColor='D9E1F2')
        cell.font = Font(name='맑은 고딕', size=10, bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[at].height = 28


def simple_sheet(book, name, headers, rows, warnings):
    sheet = book.create_sheet(name)
    header(sheet, 1, headers, warnings)
    for i, contents in enumerate(rows, 2):
        row(sheet, i, contents, warnings)
        sheet.row_dimensions[i].height = 54
    for i in range(1, len(headers)+1):
        sheet.column_dimensions[get_column_letter(i)].width = 30 if i < 4 else 45
    style_grid(sheet, 1, max(1, len(rows)+1), len(headers))
    sheet.freeze_panes = 'A2'
    return sheet


def artifact(db, p, config, workbook, filename, selection, warnings):
    buffer = BytesIO()
    workbook.save(buffer)
    content = buffer.getvalue()
    identity = new_id()
    directory = config.data_dir / 'artifacts'
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / f'{identity}.xlsx'
    target.write_bytes(content)
    try:
        saved = create(db, p, 'artifact', {'filename': filename, 'storage_path': str(target), 'content_type': XLSX,
                       'content_hash': sha256(content).hexdigest(), 'selection': selection, 'warnings': warnings,
                       'download_url': f'/api/v2/projects/{p}/artifacts/{identity}/download'}, identity)
        return make_job(db, p, 'export', selection, result_ref={'kind': 'artifact', 'id': saved.id})
    except Exception:
        target.unlink(missing_ok=True)
        raise


@router.post('/exports/traceability', status_code=202)
def trace_export(p: str, body: TraceExport, db: DB, config: Config):
    lock_project(db, p)
    run = resource(db, p, body.verification_run_id, 'verification_run')
    rows = selected_rows(run)
    if not rows:
        raise APIError(409, 'EMPTY_EXPORT', '출력할 추적표 행이 없습니다')
    if len(rows) > config.max_export_rows:
        raise APIError(413, 'EXPORT_LIMIT_EXCEEDED', '출력 행 한도를 초과했습니다')
    book, warnings = Workbook(), deepcopy(run.data['missing_features'])
    book.remove(book.active)
    selection = {**body.model_dump(), 'verification_run_version': run.version, 'snapshot_id': run.data['snapshot_id'], 'model_config_snapshot': run.data['model_config_snapshot']}
    simple_sheet(book, '요구사항 추적표', ['요구사항 ID','개정','주제','소주제','기능명','분석 상태','판정','사유','스냅샷','모델 설정'],
                 [[r['requirement_id'], r['requirement_revision'], r['topic_title'], r['subtopic_title'], r['title'],
                   r['analysis_state'], r['status'] or '미분석', r['reason'], run.data['snapshot_id'],
                   json.dumps(run.data['model_config_snapshot'], ensure_ascii=False)] for r in rows], warnings)
    condition_rows, evidence_rows, source_rows = [], [], []
    for req in run.data['requirements']:
        result = run.data['results'].get(req['id'], {})
        verdicts = {c['condition_id']: c for c in result.get('condition_verdicts', [])}
        for c in req['conditions']:
            v = verdicts.get(c['id'], {})
            condition_rows.append([req['id'], req['revision'], c['id'], c['text'], c['category'], c['target'], c['verifiable'],
                                   v.get('status') or '미분석', v.get('reason', ''), json.dumps(c['source_refs'], ensure_ascii=False)])
        for e in result.get('evidence', []):
            evidence_rows.append([req['id'], e['id'], e['chunk_id'], e['file_path'], e['start_line'], e['end_line'],
                                  e['role'], e['scope'], e['code_snippet'], run.data['snapshot_id']])
        doc = resource(db, p, req['document_id'], 'document')
        blocks = {b['id']: b for b in doc.data['blocks']}
        for ref in req['source_refs']:
            b = blocks[ref['block_id']]
            source_rows.append([req['id'], req['revision'], req['document_id'], req['document_revision'], b['id'],
                                json.dumps(b['location'], ensure_ascii=False), ref['quote']])
    if any(len(items) > config.max_export_rows for items in [condition_rows, evidence_rows, source_rows]):
        raise APIError(413, 'EXPORT_LIMIT_EXCEEDED', '추적표 상세 시트의 출력 행 한도를 초과했습니다')
    simple_sheet(book, '수락조건', ['요구사항 ID','개정','조건 ID','조건','분류','대상','검증 가능','판정','사유','출처'], condition_rows, warnings)
    simple_sheet(book, '코드 근거', ['요구사항 ID','근거 ID','청크 ID','파일','시작 줄','끝 줄','역할','범위','코드','스냅샷'], evidence_rows, warnings)
    simple_sheet(book, '원문 출처', ['요구사항 ID','개정','문서 ID','문서 개정','블록 ID','위치','원문'], source_rows, warnings)
    return artifact(db, p, config, book, 'traceability.xlsx', selection, warnings)


def date_value(text):
    return datetime.fromisoformat(text.replace('Z', '+00:00')).replace(tzinfo=None) if text else None


def test_sheet(book, name, project_name, created_at, test_date, warnings):
    sheet = book.create_sheet(name)
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 32
    row(sheet, 1, ['프로젝트명', project_name, None, None, None, None, '사양서작성자', None, '테스트확인자', None], warnings)
    row(sheet, 2, ['시험대상 메뉴명', None, name.split('. ', 1)[-1], None, None, None, '사양서작성일', date_value(created_at), '테스트시행일', date_value(test_date)], warnings)
    header(sheet, 4, ['시험번호','시험구분',None,'No','시험절차','예상결과','시험데이터',None,None,'확인결과'], warnings)
    header(sheet, 5, ['시나리오명','DEPTH(1)','DEPTH(2)',None,None,None,'계정','입력값','비고',None], warnings)
    sheet.row_dimensions[4].height = 40
    sheet.row_dimensions[5].height = 30
    for merged in ['B1:F1', 'A2:B2', 'C2:F2', 'B4:C4', 'G4:I4', 'D4:D5', 'E4:E5', 'F4:F5', 'J4:J5']:
        sheet.merge_cells(merged)
    row(sheet, 6, [None,None,None,'안내','한 행은 시험 단계 한 개입니다.','명세 기준 기대 결과',None,None,None,'정상:O / 비정상:X / 보류'], warnings)
    sheet.row_dimensions[6].height = 40
    for k, width in zip('ABCDEFGHIJ', [30,18,24,6,52,52,20,36,36,18]):
        sheet.column_dimensions[k].width = width
    sheet.freeze_panes = 'E7'
    sheet.print_title_rows = '1:6'
    return sheet


@router.post('/exports/test-specification', status_code=202)
def test_export(p: str, body: TestExport, db: DB, config: Config):
    lock_project(db, p)
    owner = project(db, p)
    spec = resource(db, p, body.test_specification_id, 'test_specification')
    spec_data = historical(db, spec, body.revision)
    cases = current_cases(db, p, spec, body.revision)
    if not cases:
        raise APIError(409, 'EMPTY_EXPORT', '선택 revision에 출력할 시험이 없습니다')
    if sum(len(c['steps']) for c in cases) > config.max_export_rows:
        raise APIError(413, 'EXPORT_LIMIT_EXCEEDED', '출력 단계 행 한도를 초과했습니다')
    run = resource(db, p, spec_data['verification_run_id'], 'verification_run')
    groups = {r['subtopic_id']: (r['topic_title'], r['subtopic_title']) for r in spec_data['requirements']}
    selected, execution_refs = {}, {}
    for case in cases:
        e = latest_execution(db, p, case)
        selected[case['id']] = deepcopy(e.data) if e else None
        execution_refs[case['id']] = {'execution_id': e.id, 'version': e.data['version'], 'test_case_revision': case['revision']} if e else None
    book, warnings = Workbook(), []
    book.remove(book.active)
    history_rows = [[r['revision'], date_value(r['date']), r['reason']] for r in spec_data['revision_history']]
    simple_sheet(book, '개정이력', ['개정','작성일자','내용'], history_rows, warnings)
    index = book.create_sheet('00. 목차')
    row(index, 1, ['목차'], warnings)
    header(index, 2, ['메뉴명','목차구분',None,None,None,'목차번호','작성여부','작성일자'], warnings)
    header(index, 3, [None,'DEPTH(1)','DEPTH(2)','DEPTH(3)','DEPTH(4)',None,None,None], warnings)
    for span in ['A2:A3','B2:E2','F2:F3','G2:G3','H2:H3']:
        index.merge_cells(span)
    row(index, 4, [None,None,None,None,None,None,'O / X',None], warnings)
    for col in 'ABCDEFGH':
        index.column_dimensions[col].width = 24
    for section, name in EXPORT_SECTIONS.items():
        relevant = [c for c in cases if c['export_section_key'] == section]
        if not relevant:
            continue
        performed = [r['performed_at'] for c in relevant if selected[c['id']] for r in selected[c['id']]['step_results'] if r.get('performed_at')]
        sheet = test_sheet(book, name, owner.name, spec.created_at, max(performed) if performed else None, warnings)
        row_number = 7
        for case in relevant:
            topic, subtopic = groups[case['subtopic_id']]
            executed = selected[case['id']]
            results = {r['step_key']: r for r in executed['step_results']} if executed else {}
            for step in case['steps']:
                actual = results.get(step['step_key'], {})
                confirmation = {'passed': 'O', 'failed': 'X', 'on_hold': '보류'}.get(actual.get('status'))
                row(sheet, row_number, [case['test_no'], topic, subtopic, step['order'], step['action'], step['expected'],
                    step['test_data']['account'] or '-', step['test_data']['input'], step['test_data']['notes'], confirmation], warnings)
                # Size by visible text/column width; no arbitrary small height for long steps.
                widths = {5:52,6:52,8:36,9:36}
                line_count = max(2, *(sum(max(1, math.ceil(len(line)*1.8/width)) for line in str(sheet.cell(row_number, col).value or '').split('\n')) for col, width in widths.items()))
                sheet.row_dimensions[row_number].height = min(409, line_count * 15 + 10)
                row_number += 1
        style_grid(sheet, 4, row_number-1, 10)
        sheet.freeze_panes = 'E7'
        paths = sorted({groups[c['subtopic_id']] for c in relevant})
        for i, (topic, subtopic) in enumerate(paths, 1):
            no = index.max_row + 1
            menu_code = name.split('.')[0]
            row(index, no, [name, topic, subtopic, None, None, f'{menu_code}-{i:03}', 'O', date_value(spec.created_at)], warnings)
            index.cell(no, 1).hyperlink = f"#'{name}'!A1"
    style_grid(index, 2, index.max_row, 8)
    index.freeze_panes = 'D5'
    if body.include_analysis_detail:
        from testgen.design import design_review
        review = design_review(spec_data, cases)
        if max(len(review['items']), len(review['conditions'])) > config.max_export_rows:
            raise APIError(413, 'EXPORT_LIMIT_EXCEEDED', '시험 설계 점검 시트의 출력 행 한도를 초과했습니다')
        statuses = {'written': '시험 연결됨', 'review_needed': '검토 필요', 'missing': '관점 연결 없음'}
        design_rows = [[i['requirement_title'], i['requirement_revision'], i['condition_text'], i['label'],
                        i['reason'], statuses[i['status']],
                        '\n'.join(f"{ref['test_no']} 단계 {ref['step_order']}" for ref in i['test_refs']),
                        i['gap_reason']] for i in review['items']]
        if review['plan_available']:
            simple_sheet(book, '시험 설계 점검', ['요구사항','개정','수락 조건','권장 시험 관점','관점 선정 기준·규칙 제안',
                         '설계 기록 연결 상태·수행 결과 아님','연결된 시험','미포함 사유'], design_rows, warnings)
        else:
            simple_sheet(book, '원문 연결 점검', ['요구사항','개정','수락 조건','연결된 시험','설계 기록'],
                         [[i['requirement_title'], i['requirement_revision'], i['condition_text'],
                           '\n'.join(f"{ref['test_no']} 단계 {ref['step_order']}" for ref in i['test_refs']),
                           '작성 당시 관점 계획과 사유 미기록'] for i in review['conditions']], warnings)
        details = []
        for case in cases:
            e = selected[case['id']]
            results = {r['step_key']: r for r in e['step_results']} if e else {}
            for step in case['steps']:
                r = results.get(step['step_key'], {})
                reqs = [next(x for x in spec_data['requirements'] if x['id'] == link['requirement_id']) for link in step['requirement_links']]
                design = step.get('design') or {}
                case_design = case.get('design') or {}
                design_status = ('편집 전 설명·재검토 필요' if case_design and case_design.get('generated_revision') != case['revision']
                                 else '해석 확인 필요' if design.get('basis') == 'review_needed'
                                 else 'AI 설계 설명·원문 대조 필요' if design else '설계 사유 미기록')
                details.append([case['test_no'], step['order'], case['revision'], json.dumps(step['requirement_links'],ensure_ascii=False),
                                json.dumps(case['prediction'],ensure_ascii=False), '\n'.join(x['original_text'] for x in reqs),
                                json.dumps([run.data['results'].get(x['id'], {}).get('evidence', []) for x in reqs],ensure_ascii=False),
                                e['attempt_no'] if e else None, r.get('status','not_run'), r.get('actual',''), r.get('notes',''),
                                r.get('performed_at'), e['environment'] if e else '', e['tested_build'] if e else '',
                                case_design.get('rationale', ''), design.get('rationale', ''), design.get('input_rationale', ''), design_status])
        simple_sheet(book, '분석 상세', ['시험번호','단계','절차 개정','조건 연결','코드 예측','원문','코드 근거','회차','수행 상태','실제 결과','수행 메모','수행 시각','환경','실제 빌드',
                     '시험 설계 이유','절차 선정 이유','입력값·역할 선정 이유','설계 설명 상태'], details, warnings)
    selection = {**body.model_dump(), 'verification_run_version': run.version, 'snapshot_id': run.data['snapshot_id'], 'case_versions': spec_data['case_versions'], 'executions': execution_refs}
    return artifact(db, p, config, book, 'test-specification.xlsx', selection, warnings)


@router.get('/artifacts/{a}/download')
def download(p: str, a: str, db: DB, config: Config):
    item = resource(db, p, a, 'artifact')
    target = Path(item.data['storage_path']).resolve()
    if not target.is_relative_to((config.data_dir / 'artifacts').resolve()) or not target.is_file():
        raise APIError(404, 'ARTIFACT_NOT_FOUND', '산출물 파일이 없거나 보관 기간이 종료되었습니다')
    return FileResponse(target, media_type=XLSX, filename=item.data['filename'])
