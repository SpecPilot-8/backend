"""골든셋(사람이 확인한 정답)을 만들고 판정을 채점한다. LLM 호출 없음.

    python scripts/golden.py worksheet            # output/golden_worksheet_<날짜>.xlsx
    python scripts/golden.py worksheet --all      # 흔들린 조건만이 아니라 전체 조건
    python scripts/golden.py import output/golden_worksheet_<날짜>.xlsx
    python scripts/golden.py score                # 레포별 최근 전체 판정을 정답과 비교

worksheet는 기본으로 "최근 두 번의 전체 판정에서 결과가 바뀐 조건"만 담는다. LLM이 가장
헷갈리는 조건이라 정답을 먼저 매길 가치가 크다. 사람은 샘플 앱을 실행해 확인하고 노란 "정답"
칸에서 상태를 고르면 된다. 판정 기준은 docs/label-definition.md.
"""

import argparse
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from matcher.data import connect
from scripts.export_trace import (BORDER, FONT, INPUT_FILL, REPOS, STATUS_FILL, STATUS_LABEL, STATUSES,
                                  _header, _row)

# 확인할 화면 주소 (진입 경로). 샘플 앱의 라우트 기준
SCREEN_PATH = {
    "LOGIN-001": "/login", "LOGIN-002": "/signup", "LOGIN-003": "/password/forgot",
    "LOGIN-004": "/password/reset (비밀번호 찾기에서 이메일·이름 확인 후 이동)",
    "NOTICE-001": "/notices", "NOTICE-002": "/notices/new (관리자)", "NOTICE-003": "/notices/new → 취소 (관리자)",
    "NOTICE-004": "/notices/{id}/edit (관리자)", "NOTICE-005": "/notices/{id}/edit → 삭제 (관리자)",
    "NOTICE-006": "/notices/{id}", "MYPAGE-001": "/mypage", "MYPAGE-002": "/mypage → 비밀번호 변경",
}
BASE_URL = {"spring-sample": "http://localhost:8080", "react-sample": "http://localhost:5173"}

GUIDE = [
    ("골든셋이란", "SpecPilot 판정을 채점하기 위한 답안지. 샘플 앱을 직접 실행해 각 조건이 실제로 구현됐는지 확인하고 정답을 고른다."),
    ("하는 일", "'확인할 조건' 시트에서 행마다 샘플 앱의 해당 화면을 열어 조건을 확인하고, 노란 '정답' 칸에서 상태를 고른다. 판단 근거는 '메모'에 한 줄."),
    ("판정 A/B", "같은 조건을 LLM이 두 번 판정한 결과. 서로 달라서 이 목록에 올라왔다. 정답을 고를 때 참고만 하고 따라가지 말 것."),
    ("판정 기준", "docs/label-definition.md (잠정). 예: 에러 문구만 다르면 implemented + 메모에 '문구 다름'. 버튼만 있고 눌러도 동작이 없으면 not_found."),
    ("needs_review 금지", "정답은 사람이 확인한 사실이라 needs_review를 쓰지 않는다. 기능이 기획과 다른 화면에 있으면 mismatch, 이 화면에 없으면 not_found. 확인이 불가능하면 비워 두고 메모에 이유."),
    ("계정", "관리자 admin@admin.com / admin1234!    일반 사용자 user@test.com / user1234!  (두 스택 동일)"),
    ("Spring 실행", "JDK 21 필요 (brew install openjdk@21). specpilot_projects_share/specpilot 에서 ./gradlew bootRun → http://localhost:8080"),
    ("React 실행", "specpilot_projects_share/specpilot-react 에서 npm install && npm install --prefix client && npm install --prefix server, 그다음 npm run dev → http://localhost:5173"),
    ("다 채운 뒤", "python scripts/golden.py import <이 파일>  → python scripts/golden.py score"),
]


def _full_runs(conn, repo_id: str) -> list[int]:
    """요구사항 전체를 판정한 실행들, 최근 것부터."""
    total = conn.execute("SELECT count(*) FROM requirement").fetchone()[0]
    return [r for (r,) in conn.execute("""
        SELECT r.id FROM verdict_run r JOIN snapshot s ON s.id = r.snapshot_id
        WHERE s.repo_id = %s
          AND (SELECT count(*) FROM requirement_verdict v WHERE v.verdict_run_id = r.id) = %s
        ORDER BY r.id DESC""", (repo_id, total))]


def worksheet(args) -> None:
    out = Path(args.out) if args.out else Path("output") / f"golden_worksheet_{date.today().isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    guide = wb.active
    guide.title = "안내"
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 110
    guide.cell(row=1, column=1, value="골든셋 작성 안내").font = Font(name=FONT, bold=True, size=14)
    for i, (k, v) in enumerate(GUIDE, start=3):
        _row(guide, i, [k, v])
        guide.cell(row=i, column=1).font = Font(name=FONT, bold=True)
    r = len(GUIDE) + 4
    _header(guide, r, ["상태", "뜻"], [16, 110])
    for j, s in enumerate(STATUSES, start=1):
        _row(guide, r + j, [s, STATUS_LABEL[s]])
        guide.cell(row=r + j, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[s])

    ws = wb.create_sheet("확인할 조건")
    headers = ["레포", "화면ID", "화면명", "확인할 주소", "요구사항", "조건", "기획서 원문",
               "판정 A", "판정 A 이유", "판정 B", "판정 B 이유", "정답", "메모", "condition_id"]
    _header(ws, 1, headers, [9, 11, 14, 32, 13, 45, 50, 14, 45, 14, 45, 16, 30, 11])

    row = 2
    with connect() as conn:
        for repo, label in REPOS:
            runs = _full_runs(conn, repo)
            if len(runs) < 2 and not args.all:
                raise SystemExit(f"{repo}: 비교할 전체 판정이 두 번 이상 없다. --all로 전체 조건을 뽑을 것")
            new, old = runs[0], (runs[1] if len(runs) > 1 else runs[0])
            labeled = {c for (c,) in conn.execute("SELECT condition_id FROM golden_label WHERE repo_id = %s", (repo,))}
            rows = conn.execute("""
                SELECT s.screen_id, s.screen_name, q.stable_key, rc.statement, q.body,
                       a.status, a.reasoning, b.status, b.reasoning, rc.id
                FROM condition_verdict b
                JOIN condition_verdict a ON a.condition_id = b.condition_id AND a.verdict_run_id = %s
                JOIN requirement_condition rc ON rc.id = b.condition_id
                JOIN requirement q ON q.id = rc.requirement_id
                JOIN screen s ON s.screen_id = q.screen_id
                WHERE b.verdict_run_id = %s AND (%s OR a.status <> b.status)
                ORDER BY s.screen_id, q.id, rc.seq""", (old, new, args.all)).fetchall()
            for sid, sname, key, stmt, body, sa, ra, sb, rb, cid in rows:
                if cid in labeled:
                    continue  # 이미 정답이 있는 조건은 다시 묻지 않는다
                url = BASE_URL[repo] + SCREEN_PATH.get(sid, "")
                _row(ws, row, [label, sid, sname, url, key, stmt, body, sa, ra, sb, rb, None, None, cid])
                for col, st in ((8, sa), (10, sb)):
                    if st in STATUS_FILL:
                        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=STATUS_FILL[st])
                ws.cell(row=row, column=12).fill = INPUT_FILL
                row += 1

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    if row > 2:
        dv.add(f"L2:L{row - 1}")
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:N{row - 1}"
    wb.save(out)
    print(f"{out} 저장: 확인할 조건 {row - 2}개")


def import_(args) -> None:
    wb = load_workbook(args.file)
    ws = wb["확인할 조건"]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}
    repo_of = {label: repo for repo, label in REPOS}
    n = skipped = 0
    with connect() as conn:
        for cells in ws.iter_rows(min_row=2, values_only=True):
            status, cid = cells[col["정답"]], cells[col["condition_id"]]
            if not status or cid is None:
                continue
            if status not in STATUSES:
                print(f"  건너뜀: condition {cid} 정답 {status!r}는 상태 목록에 없다")
                skipped += 1
                continue
            conn.execute("""
                INSERT INTO golden_label (repo_id, condition_id, status, note, source)
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT (repo_id, condition_id) DO UPDATE
                SET status = EXCLUDED.status, note = EXCLUDED.note, source = EXCLUDED.source, labeled_at = now()""",
                (repo_of[cells[col["레포"]]], cid, status, cells[col["메모"]], Path(args.file).name))
            n += 1
        conn.commit()
    print(f"정답 {n}개 저장" + (f", {skipped}개 건너뜀" if skipped else ""))


def score(args) -> None:
    with connect() as conn:
        for repo, label in REPOS:
            runs = _full_runs(conn, repo)
            if not runs:
                print(f"{label}: 전체 판정 없음")
                continue
            run = runs[0]
            pairs = conn.execute("""
                SELECT g.status, v.status, q.stable_key, rc.statement
                FROM golden_label g
                JOIN condition_verdict v ON v.condition_id = g.condition_id AND v.verdict_run_id = %s
                JOIN requirement_condition rc ON rc.id = g.condition_id
                JOIN requirement q ON q.id = rc.requirement_id
                WHERE g.repo_id = %s""", (run, repo)).fetchall()
            if not pairs:
                print(f"{label} (판정 실행 {run}): 정답이 아직 없다")
                continue
            hit = sum(g == v for g, v, _, _ in pairs)
            print(f"\n{label} (판정 실행 {run}): 정답 {len(pairs)}개 중 {hit}개 일치 = {hit / len(pairs):.0%}")

            # 상태별 재현율: 정답이 X인 것 중 판정도 X인 비율. 특히 mismatch·not_found를 놓치지 않는지가 중요하다 (9장)
            by_gold = Counter(g for g, _, _, _ in pairs)
            for st in STATUSES:
                if by_gold[st]:
                    ok = sum(1 for g, v, _, _ in pairs if g == st and v == st)
                    print(f"  정답 {st:<26} {by_gold[st]:>3}개 → 판정도 같음 {ok}개")
            wrong = [(k, s, g, v) for g, v, k, s in pairs if g != v]
            if wrong:
                print("  틀린 판정:")
                for k, s, g, v in wrong:
                    print(f"    {k} 정답 {g} / 판정 {v} — {s[:50]}")


def main() -> None:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    w = sub.add_parser("worksheet")
    w.add_argument("--all", action="store_true", help="흔들린 조건만이 아니라 전체 조건")
    w.add_argument("--out")
    i = sub.add_parser("import")
    i.add_argument("file")
    sub.add_parser("score")
    args = ap.parse_args()
    {"worksheet": worksheet, "import": import_, "score": score}[args.cmd](args)


if __name__ == "__main__":
    main()
