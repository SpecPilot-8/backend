"""판정 결과를 요구사항 추적표(xlsx)로 내보낸다 (LLM 호출 없음).

    python scripts/export_trace.py                       # output/traceability_<날짜>.xlsx
    python scripts/export_trace.py --out 추적표.xlsx

레포마다 요구사항 전체를 판정한 가장 최근 실행을 쓴다 (일부 화면만 돌린 실행은 건너뛴다).
조건 시트의 "정답" 칸은 골든셋을 만들 때 사람이 채우는 곳이다.

결과 파일에는 기획서 원문과 샘플 코드 경로가 들어가므로 output/은 git에서 제외한다.
"""

import argparse
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from matcher.data import connect

REPOS = [("spring-sample", "Spring"), ("react-sample", "React")]

STATUSES = ["implemented", "partial", "mismatch", "not_found", "needs_review", "not_statically_verifiable"]
STATUS_LABEL = {
    "implemented": "구현됨",
    "partial": "일부 구현",
    "mismatch": "불일치 (값·동작·위치가 기획서와 다름)",
    "not_found": "구현 없음",
    "needs_review": "사람 검수 필요",
    "not_statically_verifiable": "정적 검증 불가 (판정 안 함)",
}
STATUS_FILL = {
    "implemented": "E2EFDA",
    "partial": "FFF2CC",
    "mismatch": "F8CBAD",
    "not_found": "F4B084",
    "needs_review": "DDEBF7",
    "not_statically_verifiable": "EDEDED",
}
REASON_LABEL = {
    "llm": "LLM 판정",
    "llm_invalid_output": "LLM 출력 검증 실패",
    "rule_not_verifiable": "규칙: 정적 검증 불가",
    "rule_match_invalid": "규칙: 매칭 검증 실패",
    "rule_no_evidence": "규칙: 이 화면 근거 없음",
    "rule_other_screen_only": "규칙: 다른 화면 코드만 있음",
}

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")  # 사람이 채우는 칸
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _latest_full_run(conn, repo_id: str) -> int:
    """요구사항 전체를 판정한 가장 최근 실행."""
    total = conn.execute("SELECT count(*) FROM requirement").fetchone()[0]
    row = conn.execute("""
        SELECT r.id FROM verdict_run r JOIN snapshot s ON s.id = r.snapshot_id
        WHERE s.repo_id = %s
          AND (SELECT count(*) FROM requirement_verdict v WHERE v.verdict_run_id = r.id) = %s
        ORDER BY r.id DESC LIMIT 1""", (repo_id, total)).fetchone()
    if row is None:
        raise SystemExit(f"{repo_id}: 요구사항 전체를 판정한 실행이 없다. scripts/run_verdict.py를 먼저 돌릴 것")
    return row[0]


def _header(ws, row: int, headers: list[str], widths: list[int]) -> None:
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = w


def _row(ws, row: int, values: list, status_col: int | None = None) -> None:
    for col, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = Font(name=FONT)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BORDER
    if status_col is not None:
        status = values[status_col - 1]
        if status in STATUS_FILL:
            ws.cell(row=row, column=status_col).fill = PatternFill("solid", fgColor=STATUS_FILL[status])


def _condition_sheet(wb, conn, run_id: int, title: str) -> None:
    ws = wb.create_sheet(title)
    headers = ["화면ID", "화면명", "요구사항", "조건#", "조건", "정적 검증", "판정", "판정 방식", "판정 이유",
               "문구 일치", "서버 검증", "기획서 의심", "근거 코드 (이 화면)", "참고 코드 (다른 화면)",
               "정답 (골든셋)", "메모"]
    widths = [12, 16, 14, 7, 45, 10, 16, 18, 60, 10, 10, 30, 50, 40, 16, 25]
    _header(ws, 1, headers, widths)

    rows = conn.execute("""
        SELECT s.screen_id, s.screen_name, q.stable_key, rc.seq, rc.statement, rc.verifiable,
               cv.status, cv.reason_kind, cv.reasoning, cv.message_match, cv.server_validation, cv.spec_suspect,
               (SELECT string_agg(c.file_path || ':' || c.start_line || '-' || c.end_line, E'\\n' ORDER BY e.rank)
                  FROM condition_verdict_evidence e JOIN chunk c ON c.id = e.chunk_id
                 WHERE e.condition_verdict_id = cv.id AND e.role = 'primary'),
               (SELECT string_agg(c.file_path || ':' || c.start_line || '-' || c.end_line, E'\\n' ORDER BY e.rank)
                  FROM condition_verdict_evidence e JOIN chunk c ON c.id = e.chunk_id
                 WHERE e.condition_verdict_id = cv.id AND e.role = 'reference')
        FROM condition_verdict cv
        JOIN requirement_condition rc ON rc.id = cv.condition_id
        JOIN requirement q ON q.id = rc.requirement_id
        JOIN screen s ON s.screen_id = q.screen_id
        WHERE cv.verdict_run_id = %s
        ORDER BY s.screen_id, q.id, rc.seq""", (run_id,)).fetchall()

    na = {"not_applicable": "", None: ""}
    for i, r in enumerate(rows, start=2):
        (sid, sname, key, seq, statement, verifiable, status, reason_kind, reasoning,
         msg, server, suspect, primary, reference) = r
        _row(ws, i, [
            sid, sname, key, seq, statement,
            "불가" if verifiable == "not_statically_verifiable" else "가능",
            status, REASON_LABEL.get(reason_kind, reason_kind), reasoning,
            {"match": "일치", "differs": "다름"}.get(msg, na.get(msg, msg)),
            {"present": "있음", "missing": "없음"}.get(server, na.get(server, server)),
            suspect or "", primary or "", reference or "", None, None,
        ], status_col=7)
        ws.cell(row=i, column=15).fill = INPUT_FILL

    # 정답 칸은 판정 상태 중에서 고르게 한다 (라벨 정의서와 같은 enum)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    if rows:
        dv.add(f"O2:O{len(rows) + 1}")
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def _requirement_sheet(wb, conn, runs: dict[str, int]) -> None:
    ws = wb.create_sheet("요구사항")
    headers = ["화면ID", "화면명", "요구사항", "원문"] + [f"{label} 판정" for _, label in REPOS]
    _header(ws, 1, headers, [12, 16, 14, 70, 16, 16])
    rows = conn.execute("""
        SELECT s.screen_id, s.screen_name, q.stable_key, q.body,
               max(CASE WHEN v.verdict_run_id = %s THEN v.status END),
               max(CASE WHEN v.verdict_run_id = %s THEN v.status END)
        FROM requirement q JOIN screen s ON s.screen_id = q.screen_id
        LEFT JOIN requirement_verdict v ON v.requirement_id = q.id AND v.verdict_run_id IN (%s, %s)
        GROUP BY s.screen_id, s.screen_name, q.id, q.stable_key, q.body
        ORDER BY s.screen_id, q.id""",
        (runs["spring-sample"], runs["react-sample"], runs["spring-sample"], runs["react-sample"])).fetchall()
    for i, r in enumerate(rows, start=2):
        _row(ws, i, list(r))
        for col in (5, 6):
            status = r[col - 1]
            if status in STATUS_FILL:
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"


def _summary_sheet(wb, conn, runs: dict[str, int]) -> None:
    ws = wb.active
    ws.title = "요약"
    ws.column_dimensions["A"].width = 28
    for col in "BCDE":
        ws.column_dimensions[col].width = 16
    title = ws.cell(row=1, column=1, value="SpecPilot 요구사항 추적표")
    title.font = Font(name=FONT, bold=True, size=14)
    ws.cell(row=2, column=1, value=f"생성일 {date.today().isoformat()}").font = Font(name=FONT, color="808080")

    # 상태별 건수: 요구사항 시트와 조건 시트를 세는 수식. 시트를 고치면 다시 계산된다
    _header(ws, 4, ["상태", "요구사항 Spring", "요구사항 React", "조건 Spring", "조건 React"], [28, 16, 16, 16, 16])
    for i, status in enumerate(STATUSES, start=5):
        ws.cell(row=i, column=1, value=status)
        ws.cell(row=i, column=2, value=f"=COUNTIF('요구사항'!E:E,$A{i})")
        ws.cell(row=i, column=3, value=f"=COUNTIF('요구사항'!F:F,$A{i})")
        ws.cell(row=i, column=4, value=f"=COUNTIF('조건_Spring'!G:G,$A{i})")
        ws.cell(row=i, column=5, value=f"=COUNTIF('조건_React'!G:G,$A{i})")
        for col in range(1, 6):
            c = ws.cell(row=i, column=col)
            c.font = Font(name=FONT)
            c.border = BORDER
        ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
    total_row = 5 + len(STATUSES)
    ws.cell(row=total_row, column=1, value="합계").font = Font(name=FONT, bold=True)
    for col in "BCDE":
        c = ws[f"{col}{total_row}"]
        c.value = f"=SUM({col}5:{col}{total_row - 1})"
        c.font = Font(name=FONT, bold=True)

    # 실행 정보: 어느 실행 결과인지 남긴다 (원칙2: 판정은 특정 시점에 고정)
    r = total_row + 2
    _header(ws, r, ["실행 정보", "Spring", "React"], [28, 16, 16])
    info = {}
    for repo, _ in REPOS:
        info[repo] = conn.execute("""
            SELECT r.id, r.snapshot_id, r.model, r.prompt_version, r.match_run_id, r.decompose_run_id,
                   to_char(r.created_at, 'YYYY-MM-DD HH24:MI')
            FROM verdict_run r WHERE r.id = %s""", (runs[repo],)).fetchone()
    labels = ["판정 실행 번호", "스냅샷", "모델", "판정 프롬프트", "매칭 실행 번호", "조건 분해 실행 번호", "판정 시각"]
    for j, label in enumerate(labels):
        _row(ws, r + 1 + j, [label, str(info["spring-sample"][j]), str(info["react-sample"][j])])

    # 상태 설명 (기준: docs/label-definition.md)
    r = r + 2 + len(labels)
    _header(ws, r, ["상태", "뜻"], [28, 16])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for j, status in enumerate(STATUSES, start=1):
        _row(ws, r + j, [status, STATUS_LABEL[status]])
        ws.merge_cells(start_row=r + j, start_column=2, end_row=r + j, end_column=5)
        ws.cell(row=r + j, column=1).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
    note = r + len(STATUSES) + 2
    ws.cell(row=note, column=1, value=(
        "판정 기준: docs/label-definition.md (잠정). 판정은 코드를 읽고 내린 예측이며 실제 실행 결과가 아니다. "
        "조건 시트의 노란 '정답' 칸은 샘플 앱을 실행해 확인한 결과를 골든셋으로 채우는 곳이다."
    )).font = Font(name=FONT, color="808080")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", help="출력 경로 (기본: output/traceability_<날짜>.xlsx)")
    args = ap.parse_args()
    out = Path(args.out) if args.out else Path("output") / f"traceability_{date.today().isoformat()}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)

    with connect() as conn:
        runs = {repo: _latest_full_run(conn, repo) for repo, _ in REPOS}
        wb = Workbook()
        _summary_sheet(wb, conn, runs)
        _requirement_sheet(wb, conn, runs)
        for repo, label in REPOS:
            _condition_sheet(wb, conn, runs[repo], f"조건_{label}")
    # LibreOffice 없이 만들면 수식 값이 비어 있다. 엑셀이 열 때 다시 계산하게 한다
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print(f"{out} 저장 (판정 실행: " + ", ".join(f"{r} {i}" for r, i in runs.items()) + ")")


if __name__ == "__main__":
    main()
